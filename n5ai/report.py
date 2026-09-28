"""Daily Excel workbook: the calls for every Nifty 500 stock x 3 timeframes, the evidence behind them, and - as the
last sheet - the success rate of all earlier calls."""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .catalogue import catalogue
from .common import HOR, HOR_ADJ, TF_RULES, TF_SHORT

TFS = [(tf, lab, sm, tm, cm, hold) for tf, (lab, sm, tm, cm, hold) in TF_RULES.items()]
REGIME_WORD = {1: "rising", 0: "mixed", -1: "falling"}
SHORT_VERDICT = {"Works (significant)": "works", "Works in reverse (fade it)": "works in reverse", "Weak edge": "weak edge",
                 "Weak - tends to fail": "tends to fail", "No reliable edge": "no edge", "Too few events": "rare",
                 "Survivorship illusion": "survivorship illusion"}
SURV_ILLUSION = {"Event decline (>=15% in 2 days)", "Dead-cat bounce breakdown"}
PATTERN_NOTE = {
    "Event decline (>=15% in 2 days)": "Buying after a >=15% two-day fall looked profitable only because the stocks that kept "
                                       "falling later left the index. On the full NSE universe it lagged (1-yr median -14.7% "
                                       "vs average). Not used for any call.",
    "Dead-cat bounce breakdown": "Same crash-rebound effect as 'Event decline': visible only in today's survivors. Not used.",
    "Horn bottom (weekly)": "Rare spike-low pattern (95 events); excluded from the calls as a precaution.",
}
SIG_SHORT = {"Holy Grail: ADX > 30 and pullback to 20 EMA": "Holy Grail pullback (ADX > 30, back to the 20-EMA)",
             "Three lower closes in an uptrend": "three lower closes in an uptrend",
             "Inverse head & shoulders": "inverse head & shoulders",
             "RSI(2) below 5 in an uptrend (Connors)": "RSI(2) < 5 in an uptrend",
             "RSI(2) above 95 in a downtrend": "RSI(2) > 95 in a downtrend",
             "Fibonacci 38-62% rally failure (downtrend)": "Fibonacci 38-62% rally failure",
             "Supertrend (10,3) turns up": "Supertrend buy flip", "Volatility stop flips long": "volatility-stop buy flip",
             "RSI(14) rises back above 30": "RSI(14) back above 30", "Candle: tasukigap (bullish)": "bullish tasuki-gap candle"}

# ---- styles -----------------------------------------------------------------------------------------------------
F = "Arial"
f_title = Font(name=F, size=14, bold=True, color="1F2937")
f_h2 = Font(name=F, size=11, bold=True, color="1F2937")
f_sub = Font(name=F, size=10, color="4B5563")
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFF")
f_norm = Font(name=F, size=10, color="000000")
f_in = Font(name=F, size=10, color="0000FF")
f_note = Font(name=F, size=9, italic=True, color="6B7280")
f_bold = Font(name=F, size=10, bold=True, color="000000")
f_link = Font(name=F, size=9, color="1D4ED8", underline="single")
f_big = Font(name=F, size=20, bold=True, color="1F4E79")
f_tile = Font(name=F, size=9, bold=True, color="4B5563")
fill_hdr = PatternFill("solid", fgColor="1F4E79")
fill_soft = PatternFill("solid", fgColor="EEF3F7")
fill_key = PatternFill("solid", fgColor="FFF7D6")
fill_tile = PatternFill("solid", fgColor="F1F5F9")
thin = Side(style="thin", color="D1D5DB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
TYPE_FILL = {"Strong Buy": ("15803D", "FFFFFF"), "Buy": ("DCFCE7", "14532D"), "Accumulate (buy on dips)": ("ECFCCB", "365314"),
             "Hold / No fresh trade": ("F3F4F6", "374151"), "Hold (mixed signals)": ("E0E7FF", "312E81"),
             "Avoid (likely underperformer)": ("FEE2E2", "7F1D1D"), "Reduce / Avoid": ("FEE2E2", "7F1D1D"),
             "Short Sell": ("FCA5A5", "7F1D1D"), "Strong Short Sell": ("B91C1C", "FFFFFF"), "Not rated": ("E5E7EB", "6B7280")}
VERDICT_FILL = {"Works (significant)": ("15803D", "FFFFFF"), "Weak edge": ("DCFCE7", "14532D"),
                "No reliable edge": ("F3F4F6", "374151"), "Weak - tends to fail": ("FEF3C7", "78350F"),
                "Works in reverse (fade it)": ("FEE2E2", "7F1D1D"), "Too few events": ("FFFFFF", "9CA3AF"),
                "Survivorship illusion": ("FDE68A", "78350F")}
RES_FILL = {"Supportive": ("DCFCE7", "14532D"), "Mixed": ("FEF3C7", "78350F"), "Negative": ("FEE2E2", "7F1D1D")}
STATUS_FILL = {"Target hit": ("15803D", "FFFFFF"), "Correct": ("DCFCE7", "14532D"), "Time exit": ("FEF3C7", "78350F"),
               "Stop hit": ("FCA5A5", "7F1D1D"), "Wrong": ("FEE2E2", "7F1D1D"), "Open": ("DBEAFE", "1E3A8A"),
               "Pending": ("F3F4F6", "374151"), "Not triggered": ("E5E7EB", "6B7280")}
OUTCOME_FILL = {"Win": ("DCFCE7", "14532D"), "Loss": ("FEE2E2", "7F1D1D")}
PCT, PCT_S2, RS_ = "0.0%", "+0.00%;-0.00%;0.00%", "#,##0.00"
# machine names of the 25 'Signals' columns (same order as the sheet) - used for the APEX export
SIGNAL_KEYS = ["symbol", "company", "sector", "timeframe", "ai_recommendation", "recommendation_type", "conf_buy",
               "conf_short", "chance_profit", "exp_return", "edge_pts", "current_price", "atr14", "entry_price",
               "target_price", "stop_loss", "potential_profit", "potential_loss", "max_days", "typical_days",
               "current_trend", "evidence", "chart_patterns", "results_check", "remarks"]
PTS = '+0.0" pts";-0.0" pts";0.0" pts"'


def header(ws, row, cols, height=34):
    for j, h in enumerate(cols, 1):
        c = ws.cell(row, j, h)
        c.font, c.fill, c.alignment, c.border = f_hdr, fill_hdr, center, border
    ws.row_dimensions[row].height = height


def cf_text(ws, rng, mapping):
    for val, (bg, fg) in mapping.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{val}"'],
                                                      fill=PatternFill("solid", fgColor=bg),
                                                      font=Font(name=F, size=10, bold=True, color=fg)))


def clean(v):
    if isinstance(v, (float, np.floating)):
        return None if not np.isfinite(v) else float(v)
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.bool_):
        return bool(v)
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    return v


def put(ws, r, c, v, font=f_norm, fmt=None, wrap=False, fill=None):
    cell = ws.cell(r, c, clean(v))
    cell.font, cell.border = font, border
    if fmt:
        cell.number_format = fmt
    if wrap:
        cell.alignment = wrap_top
    if fill is not None:
        cell.fill = fill
    return cell


def widths(ws, ws_widths):
    for j, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def bullets(ws, r, lines, cols=8, cpl=165):
    """Merged, wrapped bullet rows; the row height follows the text length (Excel does not auto-fit merged cells)."""
    for t in lines:
        r += 1
        c = ws.cell(r, 1, "• " + t)
        c.font, c.alignment = f_norm, wrap_top
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=cols)
        n_lines = max(1, -(-(len(t) + 2) // cpl))
        ws.row_dimensions[r].height = max(18, n_lines * 13.5 + 6)
    return r


def top_pct(q):
    return max(1, int(round(100 * (1 - q))))


def rank_txt(q):
    return f"top {top_pct(q)}%" if q >= 0.5 else f"bottom {max(1, int(round(100 * q)))}%"


def ordinal(n):
    n = int(n)
    suf = "th" if 10 <= n % 100 <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suf}"


def sig_list(d, key):
    v = d.get(key)
    return v if isinstance(v, list) else []


def evidence_text(d) -> str:
    parts = []
    for p in sig_list(d, "pos_signals"):
        parts.append(f"+ {p[0]} ({p[6]:%d-%b}): {100 * p[2]:+.1f}% vs avg stock, t {abs(p[4]):.1f}")
    for p in sig_list(d, "neg_signals"):
        tag = " - fades" if p[5].startswith("Works in reverse") else ""
        parts.append(f"- {p[0]}{tag} ({p[6]:%d-%b}): {-100 * p[2]:+.1f}% vs avg stock, t {abs(p[4]):.1f}")
    return "; ".join(parts)


def trend(x, tf):
    if tf == "ST":
        a, b = x["dist_ema20"], x["ta_SUPERTREND_trend"]
        if a > 0 and b > 0:
            return "Up"
        if a < 0 and b < 0:
            return "Down"
        return "Sideways"
    a, s = (x["dist_sma50"], x["slope_sma50"]) if tf == "MT" else (x["dist_sma200"], x["slope_sma200"])
    if a > 0 and s > 0:
        return "Up"
    if a < 0 and s < 0:
        return "Down"
    return "Sideways"


def disp_verdict(pattern: str, raw) -> str:
    raw = raw if isinstance(raw, str) else ""
    if pattern in SURV_ILLUSION and raw not in ("No reliable edge", "Too few events", ""):
        return "Survivorship illusion"
    return raw


def fmt_pct(v, digits=0, signed=False):
    if v is None or not np.isfinite(v):
        return "-"
    return f"{100 * v:+.{digits}f}%" if signed else f"{100 * v:.{digits}f}%"


def xl_round(v, digits=2):
    """ROUND() as Excel does it (half away from zero on the decimal value), so exported levels equal the workbook."""
    from decimal import ROUND_HALF_UP, Decimal
    if v is None or not np.isfinite(v):
        return None
    q = Decimal(1).scaleb(-digits)
    return float(Decimal(repr(float(v))).quantize(q, rounding=ROUND_HALF_UP))


# =====================================================================================================================
def build_workbook(ctx: dict, outfile: Path) -> Path:
    """ctx: see nifty500_daily.py (decisions, current, universe, market, events, regime, calib, static_dir, charts,
    track, research, research_note, asof)."""
    D = ctx["decisions"]
    CUR = ctx["current"].set_index("symbol")
    UNI = ctx["universe"].set_index("symbol")
    MK = ctx["market"]
    S_DIR = Path(ctx["static_dir"])
    RC = pd.read_csv(S_DIR / "report_card.csv")
    FOLD = pd.read_csv(S_DIR / "fold_metrics.csv")
    VAR = pd.read_csv(S_DIR / "model_variants.csv")
    BER = pd.read_csv(S_DIR / "bucket_excess_by_regime.csv")
    SURV = pd.read_csv(S_DIR / "survivorship.csv")
    CRASH = pd.read_csv(S_DIR / "crash_check.csv")
    feat_cols = json.load(open(S_DIR / "feature_columns.json"))
    FEATS = ctx["feats"]
    REGIME_NOW = int(ctx["regime"])
    RW = REGIME_WORD[REGIME_NOW]
    calib = ctx["calib"]
    BT = {"tabs": calib["tabs"][REGIME_NOW], "signals": calib["signals"][REGIME_NOW]}
    RESEARCH = ctx.get("research") or {}
    CHARTS = ctx.get("charts") or []
    TRACK = ctx.get("track") or {}
    asof = pd.Timestamp(ctx["asof"])
    ASOF = f"{asof:%a %d-%b-%Y}"
    dates = MK.index
    recent_from = dates[-10]
    EV = ctx["events"]
    EV = EV[EV["date"] >= recent_from].copy() if len(EV) else EV
    RCP = RC[(RC.level == "pattern")].drop_duplicates("pattern").set_index("pattern")

    def verdict_of(pattern: str, tf: str = "MT") -> str:
        if pattern in RCP.index:
            return SHORT_VERDICT.get(disp_verdict(pattern, RCP.at[pattern, f"{tf}_verdict"]), "")
        return ""

    def active_patterns(sym: str, maxn=4) -> str:
        if not len(EV):
            return ""
        e = EV[(EV["symbol"] == sym) & ~EV["family"].isin(["Range contraction"])].sort_values("date", ascending=False)
        seen, out = set(), []
        for _, r in e.iterrows():
            if r["pattern"] in seen:
                continue
            seen.add(r["pattern"])
            v = verdict_of(r["pattern"])
            out.append(f"{'▲' if r['dir'] > 0 else '▼'} {r['pattern']} {r['date']:%d-%b}" + (f" [{v}]" if v else ""))
            if len(out) >= maxn:
                break
        return "; ".join(out)

    def model_counts(tf):
        tabs = calib["tabs"][REGIME_NOW][(tf, "L")]["table"]
        return bool(tabs.iloc[4]["edge_ok"] or tabs.iloc[5]["edge_ok"])

    def remarks(sym, tf, d, x):
        if not d["eligible"]:
            bars = int(UNI.at[sym, "bars"]) if sym in UNI.index and pd.notna(UNI.at[sym, "bars"]) else 0
            if bars < 210:
                return f"Not rated: only {bars} sessions of history (the model needs 210)."
            return "Not rated: average daily turnover below the Rs 10 Cr liquidity floor."
        out = []
        typ = d["type"]
        h = HOR[tf]
        if typ == "Strong Buy":
            out.append(f"AI model ranks it in the top {top_pct(d['pct_buy'])}% of stocks for a {HOR_ADJ[tf]} buy. Historically, "
                       f"top-5% names hit the target {100 * d['conf_buy']:.0f}% of the time vs {100 * d['base_hit_buy']:.0f}% "
                       f"for the average stock (weighted to {RW} markets like now), averaging "
                       f"{100 * d['exp_ret_buy']:+.1f}% per trade after costs.")
        elif typ in ("Buy", "Accumulate (buy on dips)"):
            if d.get("long_model"):
                out.append(f"AI model ranks it in the top {top_pct(d['pct_buy'])}% for a {HOR_ADJ[tf]} buy.")
            if sig_list(d, "pos_signals"):
                p = sig_list(d, "pos_signals")[0]
                also = f"also positive in {RW} markets" if REGIME_NOW != 0 else "positive in every market type"
                out.append(f"Proven setup active: {p[0]} ({p[6]:%d-%b}) - historically {100 * p[2]:+.1f}% vs the average "
                           f"stock over {h} (t {abs(p[4]):.1f}, {also}).")
            out.append(f"Buy confidence {100 * d['conf_buy']:.0f}% vs {100 * d['base_hit_buy']:.0f}% for an average stock; "
                       f"expected {100 * d['exp_ret_buy']:+.1f}% per trade after costs.")
            if typ.startswith("Accumulate"):
                out.append("Build the position in parts on dips toward the 20/50-DMA rather than all at once.")
        elif typ in ("Short Sell", "Strong Short Sell"):
            out.append(f"AI model ranks it among the weakest stocks for {h} (short-side rank top {top_pct(d['pct_short'])}%); "
                       f"such short trades hit target {100 * d['conf_short']:.0f}% of the time historically, "
                       f"averaging {100 * d['exp_ret_short']:+.1f}% per trade after costs.")
        elif typ.startswith("Avoid") or typ.startswith("Reduce"):
            if d.get("avoid_model"):
                out.append(f"AI model flags a likely underperformer over {h} (short-side rank top {top_pct(d['pct_short'])}%): "
                           f"such stocks lagged the average stock by about {100 * max(d['model_ex_short'], 0):.1f}% per trade.")
            if sig_list(d, "neg_signals"):
                p = sig_list(d, "neg_signals")[0]
                how = ("has been followed by underperformance" if p[5].startswith("Works in reverse")
                       else "has a proven bearish edge")
                out.append(f"Warning signal: {p[0]} ({p[6]:%d-%b}) {how} ({-100 * p[2]:+.1f}% vs average over {h}).")
            out.append("Holders: tighten stops or reduce; no fresh buying." if tf != "ST" else "No fresh buying for now.")
        elif typ == "Hold (mixed signals)":
            parts = []
            if d.get("long_model"):
                parts.append(f"AI model ranks it in the top {top_pct(d['pct_buy'])}% for a {HOR_ADJ[tf]} buy")
            if d.get("avoid_model"):
                parts.append(f"AI model flags a likely underperformer (short-side rank top {top_pct(d['pct_short'])}%)")
            ev = evidence_text(d)
            if ev:
                parts.append(ev)
            out.append("Mixed evidence - " + (" | ".join(parts) or "model and signals disagree") + ". No fresh trade.")
        else:
            if d["pct_buy"] >= 0.85 and not d.get("long_model"):
                out.append(f"The model ranks it in the top {top_pct(d['pct_buy'])}% for a {HOR_ADJ[tf]} buy, but at this "
                           f"horizon its top ranks have not beaten the average stock in {RW} markets, so that is not counted; "
                           "no proven signal active. Holders can keep; no fresh entry.")
            else:
                out.append(f"No edge either way for {h} (buy-side model rank: {ordinal(max(1, round(100 * d['pct_buy'])))} "
                           "percentile; no proven signal active). Holders can keep; no fresh entry.")
        facts = []
        if pd.notna(x.get("ta_RSI")):
            facts.append(f"RSI {x['ta_RSI']:.0f}")
        if pd.notna(x.get("xs_rs_rating")):
            facts.append(f"RS rating {int(x['xs_rs_rating'])}")
        if pd.notna(x.get("dist_52w_high")):
            facts.append(f"{x['dist_52w_high']:.0f}% from 52-wk high")
        if tf == "ST":
            facts.append(f"{'above' if x['dist_ema20'] > 0 else 'below'} 20-EMA")
        elif tf == "MT":
            facts.append(f"{'above' if x['dist_sma50'] > 0 else 'below'} a "
                         f"{'rising' if x['slope_sma50'] > 0 else 'falling'} 50-DMA")
        else:
            facts.append(f"{'above' if x['dist_sma200'] > 0 else 'below'} a "
                         f"{'rising' if x['slope_sma200'] > 0 else 'falling'} 200-DMA")
        dz = x.get("deliv_pct_5_minus_20")
        if pd.notna(dz) and abs(dz) >= 5:
            facts.append(f"delivery % {'rising' if dz > 0 else 'falling'} ({dz:+.0f} pts vs 20-day)")
        if x.get("atr_pct_14", 0) > 5:
            facts.append(f"volatile ({x['atr_pct_14']:.1f}% daily range)")
        out.append("Technicals: " + "; ".join(facts) + ".")
        if sym in RESEARCH:
            r = RESEARCH[sym]
            out.append(f"Results {r['quarter']} - revenue: {r['revenue']}; profit: {r['profit']} ({r['verdict'].lower()}).")
        return " ".join(out)

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True           # formulas are calculated when Excel opens the file
    # =================================================================================================================
    # Read Me
    # =================================================================================================================
    rm = wb.active
    rm.title = "Read Me"
    n50 = MK.iloc[-1]
    reg_txt = {1: "RISING (Nifty 50 above its 50- and 200-DMA)", 0: "MIXED (Nifty 50 between its 50- and 200-DMA)",
               -1: "FALLING (Nifty 50 below its 50- and 200-DMA)"}[REGIME_NOW]
    rm["A1"] = f"Nifty 500 AI signals - daily run for {ASOF}"
    rm["A1"].font = f_title
    rm["A2"] = (f"NSE data to the close of {ASOF}. Market regime: {reg_txt}. Breadth: {100 * n50['mkt_pct_above50']:.0f}% of "
                f"Nifty 500 stocks above their 50-DMA, {100 * n50['mkt_pct_above200']:.0f}% above their 200-DMA.")
    rm["A2"].font = f_sub
    rm["A3"] = ("Research output for paper trading and education - not investment advice; not a SEBI-registered adviser. "
                "Every probability is a historical frequency, not a promise.")
    rm["A3"].font = Font(name=F, size=10, bold=True, color="B91C1C")
    hd = TRACK.get("headline", {})
    if hd.get("closed"):
        tr_txt = (f"Track record of these daily calls: {fmt_pct(hd['success_rate'])} of {hd['closed']} completed calls were "
                  f"right (since {hd['first_date']})"
                  + (f"; days whose calls have all finished: {fmt_pct(hd['success_rate_mature'])} of {hd['mature']}"
                     if hd.get("mature") else "") + ". Full detail on the last sheet, 'Success Rate'.")
    else:
        tr_txt = ("Track record: no call has completed yet (2-week calls finish after 10 sessions). The last sheet, "
                  "'Success Rate', fills in automatically as calls complete.")
    rm["A4"] = tr_txt
    rm["A4"].font = Font(name=F, size=10, bold=True, color="1F4E79")
    counts = D[D.eligible].groupby(["tf", "type"]).size()
    summary = []
    for tf, lab, *_ in TFS:
        c = counts.get(tf, pd.Series(dtype=int))
        buys = int(c.get("Strong Buy", 0) + c.get("Buy", 0) + c.get("Accumulate (buy on dips)", 0))
        avoid = int(c.get("Avoid (likely underperformer)", 0) + c.get("Reduce / Avoid", 0))
        shorts = int(c.get("Short Sell", 0) + c.get("Strong Short Sell", 0))
        summary.append(f"{lab}: {buys} buy calls ({int(c.get('Strong Buy', 0))} Strong Buy), "
                       + (f"{shorts} short sells, " if shorts else "") + f"{avoid} avoid/reduce, the rest hold.")
    r = 5
    rm.cell(r + 1, 1, "Today's calls in one line per timeframe").font = f_h2
    r = bullets(rm, r + 1, summary)
    r += 1
    rm.cell(r + 1, 1, "What goes into every run").font = f_h2
    n_events = int(RC[RC.level == "pattern"]["n"].sum())
    n_cp = int(((RC.level == "pattern") & (RC.source == "Chart pattern")).sum())
    n_sig = int(((RC.level == "pattern") & (RC.source == "Indicator signal")).sum())
    n_cdl = int(((RC.level == "pattern") & (RC.source == "Candlestick")).sum())
    r = bullets(rm, r + 1, [
        f"{len(feat_cols)} indicator values per stock per day: the complete TA-Lib library (136 functions) plus 60+ tools it "
        "lacks - Ichimoku, Supertrend, Renko, Kagi, three-line break, Heikin-Ashi, point & figure, volume profile, anchored "
        "VWAPs, pivot points, Laguerre/Connors RSI, Fisher, Schaff, KST, Klinger, Hurst, NSE delivery-% analytics, relative "
        "strength vs Nifty and sector, market breadth, calendar effects. Every value uses only past data.",
        f"{n_cp} chart-pattern types found by a swing engine at three swing sizes: head & shoulders, double/triple tops & "
        "bottoms, cup & handle, triangles, wedges, flags, pennants, rectangles, channels, broadening formations, diamonds, "
        "rounding tops/bottoms, V-reversals, islands, gaps, Darvas boxes, VCP, flat bases, high-tight flags, Wyckoff "
        "springs/upthrusts, 2B/turtle soup, Dow-theory reversals, Elliott 5-wave counts, three drives, harmonics (Gartley, Bat, "
        "Butterfly, Crab, Deep crab, Cypher, AB=CD), Wolfe waves, Quasimodo, Fibonacci pull-backs, trendline and "
        "support/resistance breaks, divergences, pipes & horns, point & figure signals.",
        f"{n_sig} classic indicator signals and {n_cdl} candlestick signals (all 61 TA-Lib patterns, raw and in trend "
        "context, plus tweezers, pin bars, outside bars).",
        f"Report card: {n_events:,} historical events on 500 stocks, 2017-2026, each measured against the same-day average "
        "stock. See 'Pattern Report Card'.",
        f"AI model: gradient-boosted decision trees over {len(FEATS)} features, one model per timeframe and side, trained "
        f"walk-forward ({calib.get('model_version', 'v2')}). See 'Model Validation'.",
    ])
    r += 1
    proven = {}
    for tf_, sigs in BT["signals"].items():
        for (pat, dr), info in sigs.items():
            kind = "buy" if info[0] > 0 else ("warn" if info[5].startswith("Works (") else "rev")
            proven.setdefault((kind, SIG_SHORT.get(pat, pat)), []).append(TF_SHORT[tf_])

    def proven_list(kind):
        return ", ".join(f"{name} [{'/'.join(tfs)}]" for (k, name), tfs in proven.items() if k == kind) or "none"

    n_short = int(D["type"].isin(["Short Sell", "Strong Short Sell"]).sum())
    buy_syms = sorted(set(D.loc[D["side"] > 0, "symbol"]))
    res_have = [s for s in buy_syms if s in RESEARCH]
    rm.cell(r + 1, 1, "How a call is made (same rules for every stock)").font = f_h2
    how = [
        "Evidence 1 - the AI model ranks every stock against the others each day. A rank bucket only counts if, out of sample, "
        f"it beat the same-day average stock BOTH over all 2020-26 markets AND in markets like today's ({RW}).",
        "Evidence 2 - proven signals: patterns / indicator signals that beat the average stock with |t| >= 3 in that "
        "timeframe, with the same sign in 2017-21 and 2022-26 and in today's market type, a meaningful size, and no opposite "
        f"result at the other horizons ([2w/3m/1y] = timeframes where each is used). Buy side: {proven_list('buy')}. "
        f"Warning side: {proven_list('warn')}; plus bullish signals that reliably worked in reverse: {proven_list('rev')}.",
        "Strong Buy = model top 5% with no warning. Buy (Accumulate for long term) = model top 15% or a proven buy setup, with "
        "no warning. Avoid / Reduce = model flags a likely underperformer or a proven warning is active. Hold = nothing "
        "reliable either way.",
        ("Short Sell appears only if short trades in that bucket made money after costs historically - "
         + (f"{n_short} qualify today." if n_short else "none qualify today. Use Avoid/Reduce as the bearish call.")),
        "AI confidence - buy side / short side: probability that a trade opened at the next open reaches its target before its "
        "stop, from historical trades in the same rank bucket and market type, corrected for survivorship bias (long side).",
        "Chance of profit: probability the trade ends with any gain after ~0.3% costs. Expected return: average result per "
        "trade. Edge: confidence minus the hit rate of an average stock in today's market.",
        "Entry: enter at the next open only up to this price (long). Target and stop: the rules table below (multiples of the "
        "14-day ATR). Days to hold: maximum sessions before a time exit.",
    ]
    if RESEARCH:
        how.append(f"Results & news: {ctx.get('research_note', '')} Available for {len(res_have)} of today's {len(buy_syms)} "
                   "buy-call stocks. Shown beside each call ('Top Picks', 'Signals') but NOT blended into the confidence %.")
    else:
        how.append("Results & news: not part of the daily job (the last manual check has expired). Check the latest "
                   "quarterly results before acting on any call.")
    r = bullets(rm, r + 1, how)
    r += 1
    rm.cell(r + 1, 1, "Trade rules per timeframe (levels are multiples of the 14-day ATR, the stock's average daily range)"
            ).font = f_h2
    r += 2
    header(rm, r, ["Timeframe", "Stop (x ATR)", "Target (x ATR)", "No-chase (x ATR)", "Max days", "Reward : risk", "Notes", ""])
    RULE_ROW = {}
    for tf, lab, sm, tm, cm, hold in TFS:
        r += 1
        RULE_ROW[tf] = r
        put(rm, r, 1, lab)
        for j, v in ((2, sm), (3, tm), (4, cm), (5, hold)):
            put(rm, r, j, v, font=f_in, fill=fill_key)
        put(rm, r, 6, f"=C{r}/B{r}", fmt='0.0"x"')
        put(rm, r, 7, "Blue cells are inputs: change them and every Entry/Target/Stop on 'Signals' recalculates. The "
                      "success-rate tracker always uses the published levels.", font=f_note, wrap=True)
        rm.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
        rm.row_dimensions[r].height = 28
    r += 1
    rm.cell(r + 1, 1, "What the research evidence says (read before trading)").font = f_h2
    v = VAR[VAR.model.str.startswith("Absolute")].set_index(["tf", "side"])
    st_b = calib["tabs"][-1][("ST", "L")]["table"].iloc[5]
    cp = RC[(RC.level == "pattern") & (RC.source == "Chart pattern")]
    works = [p for p in cp[cp["MT_verdict"] == "Works (significant)"]["pattern"] if p not in SURV_ILLUSION]
    rev = [p for p in cp[cp["MT_verdict"] == "Works in reverse (fade it)"]["pattern"] if p not in SURV_ILLUSION]
    works_dirs = {int(RCP.at[p, "dir"]) for p in works}
    works_tag = " - both bearish warnings" if works_dirs == {-1} and len(works) == 2 else ""
    classic = [("Double bottom", "double bottom"), ("Symmetrical triangle - up breakout", "symmetrical-triangle breakout"),
               ("Ascending triangle - up breakout", "ascending-triangle breakout"), ("Bull flag", "bull flag"),
               ("Cup (no handle) breakout", "cup breakout"), ("Head & shoulders top", "head & shoulders top")]

    def tfmt(t):
        return f"{(t if abs(t) >= 0.05 else 0.0):+.1f}"

    classic_txt = ", ".join(f"{lab} (t {tfmt(RCP.at[p, 'MT_xnet_t'])})" for p, lab in classic if p in RCP.index)
    bad_txt = " and ".join(f"{lab} (t {tfmt(RCP.at[p, 'MT_xnet_t'])})" for p, lab in
                           (("Inverse head & shoulders", "inverse head & shoulders"), ("Wyckoff spring", "Wyckoff springs"))
                           if p in RCP.index)
    works_txt = ", ".join(p.replace(" (downtrend)", " in a downtrend") for p in works)
    s_mt_p = SURV[(SURV.tf == "MT") & SURV.universe.str.startswith("Proxy")].iloc[0]
    s_mt_t = SURV[(SURV.tf == "MT") & SURV.universe.str.startswith("Today")].iloc[0]
    s_lt_p = SURV[(SURV.tf == "LT") & SURV.universe.str.startswith("Proxy")].iloc[0]
    s_lt_t = SURV[(SURV.tf == "LT") & SURV.universe.str.startswith("Today")].iloc[0]
    cr_p = CRASH[CRASH.universe.str.startswith("Proxy")].iloc[0]
    cr_s = CRASH[CRASH.universe.str.startswith("Surv")].iloc[0]
    mlt = [tf for tf in ("MT", "LT") if model_counts(tf)]
    mid_txt = ("Mid and long term: the model's buy ranking did NOT beat the average stock in falling markets (its favourites "
               "lagged, beaten-down stocks rebounded more)")
    mid_txt += (f"; in {RW} markets like today it did for {' and '.join(HOR[t] for t in mlt)}, so model buys are allowed "
                "there today." if mlt else ", so mid/long-term buys today come only from proven setups.")
    mid_txt += (" Its warnings (likely underperformers) held up in rising, mixed and falling markets and in 19 of 20 test "
                "years across the three horizons - the most dependable output.")
    r = bullets(rm, r + 1, [
        f"Short term is where the model is most reliable: its daily ranking beat chance in 7 of 7 years (rank correlation "
        f"{v.loc[('ST', 'L'), 'daily_ic']:.3f}); the top 5% beat the same-day average stock by "
        f"{v.loc[('ST', 'L'), 'top5_excess']:+.2f}% per 2-week trade overall and {st_b['ex_ret_regime'] * 100:+.2f}% in "
        "falling markets. Small per trade - costs and discipline matter.",
        mid_txt,
        f"Chart patterns alone: only {len(works)} of {len(cp)} pattern types beat the average stock reliably over 3 months "
        f"({works_txt}{works_tag}); {len(rev)} worked in reverse; 'buy after a crash' looked like a winner but was a "
        f"survivorship illusion (below). Textbook patterns did not beat the average stock on NSE in 2017-26: {classic_txt}; "
        f"{bad_txt} did reliably worse than average (3-month t-stats; |t| >= 3 needed).",
        "What did work: buying orderly pullbacks inside strong uptrends (Holy Grail, three lower closes, RSI(2) < 5; IBS < 0.2 "
        "and Double 7s too, with smaller effects), stocks whose relative-strength line makes a 52-week high or whose moving "
        "averages newly stack bullish, and avoiding stocks in stage-4 downtrends or with RSI(2) > 95 in a downtrend. Late "
        "trend-following buy flips (Supertrend, Parabolic SAR, volatility stop, Renko/three-line-break reversals, 20-day "
        "breakouts) tended to underperform afterwards.",
        f"Survivorship check on the full NSE universe as it stood each day: an average 3-month trade hit target "
        f"{s_mt_p['hit']:.0f}% vs {s_mt_t['hit']:.0f}% on today's constituents; 1-year {s_lt_p['hit']:.0f}% vs "
        f"{s_lt_t['hit']:.0f}%. All long-side confidences here are corrected for this.",
        f"'Buy after a 15% two-day crash' looked great on today's list ({cr_s['x250_mean']:+.1f}% vs average over a year) but "
        f"not on the full universe (median {cr_p['x250_median']:+.1f}%; {100 * cr_p['share_not_in_today_list']:.0f}% of such "
        "crashes were in stocks that later left the index) - so it is NOT used as a buy reason.",
    ])
    r += 1
    rm.cell(r + 1, 1, f"Market on {ASOF} (computed from NSE data)").font = f_h2
    r = bullets(rm, r + 1, ctx.get("market_lines", []))
    if ctx.get("data_checks"):
        r += 1
        rm.cell(r + 1, 1, "Data checks for this run").font = f_h2
        r = bullets(rm, r + 1, ctx["data_checks"])
    r += 2
    c = rm.cell(r, 1, ctx.get("data_note", ""))
    c.font, c.alignment = f_note, wrap_top
    rm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    rm.row_dimensions[r].height = 40
    widths(rm, [30, 13, 13, 15, 10, 13, 40, 40])

    # =================================================================================================================
    # Signals
    # =================================================================================================================
    sg = wb.create_sheet("Signals")
    cols = ["Symbol", "Company", "Sector", "Timeframe", "AI recommendation", "Recommendation type",
            "AI confidence - buy side", "AI confidence - short side", "Chance of profit (recommended side)",
            "Expected return per trade", "Edge vs average stock", "Current price Rs", "ATR 14 Rs",
            "Entry Rs (buy up to)", "Target Rs (sell at)", "Stop loss Rs", "Potential profit", "Potential loss",
            "Days to hold (max)", "Typical days to target", "Current trend", "Evidence (AI model rank + proven signals)",
            "Active chart patterns (last 10 sessions) [3-mth verdict]", "Latest results check", "Remarks"]
    header(sg, 1, cols, 52)
    TYPE_ORDER = {"Strong Buy": 0, "Buy": 1, "Accumulate (buy on dips)": 1, "Strong Short Sell": 1.5, "Short Sell": 1.5,
                  "Hold (mixed signals)": 2, "Hold / No fresh trade": 3, "Avoid (likely underperformer)": 4,
                  "Reduce / Avoid": 4, "Not rated": 5}
    score = D.assign(o=D["type"].map(TYPE_ORDER).fillna(3)).groupby("symbol")["o"].mean()
    sub = D[D.tf == "ST"].set_index("symbol")["pct_buy"] if "pct_buy" in D else pd.Series(dtype=float)
    uni_syms = [s for s in UNI.index if UNI.at[s, "status"] == "ok"]
    order = sorted(uni_syms, key=lambda s: (score.get(s, 9), -(sub.get(s) if pd.notna(sub.get(s)) else -1), s))
    RM = "'Read Me'!"
    Dk = {(s, t): g.iloc[0].to_dict() for (s, t), g in D.groupby(["symbol", "tf"])}
    SIG_ROWS = []                        # the same rows as plain values (levels computed like the Excel formulas)

    def sig_record(vals, tf, side, sm, tm, cm):
        rec = dict(zip(SIGNAL_KEYS, [clean(v) for v in vals]))
        rec.update(tf=tf, side=int(side), rank=len(SIG_ROWS) + 1)
        L, M = rec["current_price"], rec["atr14"]
        if side != 0 and L and M is not None and np.isfinite(M):
            sg_ = 1 if side > 0 else -1
            rec["entry_price"] = xl_round(L + sg_ * cm * M)
            rec["target_price"] = xl_round(L + sg_ * tm * M)
            rec["stop_loss"] = xl_round(L - sg_ * sm * M)
            rec["potential_profit"] = abs(rec["target_price"] - L) / L
            rec["potential_loss"] = abs(L - rec["stop_loss"]) / L
        SIG_ROWS.append(rec)

    rr = 2
    for sym in order:
        comp = UNI.at[sym, "company"]
        sect = str(UNI.at[sym, "industry"]).title().replace("It", "IT")
        x = CUR.loc[sym] if sym in CUR.index else None
        pats = active_patterns(sym) if x is not None else ""
        for tf, lab, sm, tm, cm, hold in TFS:
            d = Dk.get((sym, tf))
            if d is None or x is None:
                note = UNI.at[sym, "note"] if isinstance(UNI.at[sym, "note"], str) else "no trading data."
                vals = [sym, comp, sect, lab, "Not rated", "Not rated"] + [None] * 18 + [f"Not rated: {note}"]
                for j, val in enumerate(vals, 1):
                    put(sg, rr, j, val)
                sig_record(vals, tf, 0, sm, tm, cm)
                rr += 1
                continue
            side, rated = d["side"], d["eligible"]
            win = d.get("win_buy") if side >= 0 else d.get("win_short")
            exp = d.get("exp_ret_buy") if side >= 0 else d.get("exp_ret_short")
            edge = ((d.get("conf_buy", np.nan) - d["base_hit_buy"]) if side >= 0
                    else (d.get("conf_short", np.nan) - d["base_hit_short"]))
            model_txt = ""
            if rated:
                model_txt = (f"Model rank: {rank_txt(d['pct_buy'])} buy side / {rank_txt(d['pct_short'])} short side"
                             f" (buy bucket: {d.get('bucket_buy', '')})")
            ev_txt = "; ".join(t for t in (model_txt, evidence_text(d)) if t)
            vals = [sym, comp, sect, lab, d["rec"], d["type"],
                    d.get("conf_buy") if rated else None, d.get("conf_short") if rated else None,
                    win if rated else None, exp if rated else None, (100 * edge) if rated else None,
                    round(d["close"], 2), float(d["atr"]), None, None, None, None, None,
                    hold if side != 0 else None, d.get("med_days") if side != 0 else None,
                    trend(x, tf), ev_txt, pats, RESEARCH[sym]["verdict"] if sym in RESEARCH else "Not checked",
                    remarks(sym, tf, d, x)]
            for j, val in enumerate(vals, 1):
                put(sg, rr, j, val)
            sig_record(vals, tf, side, sm, tm, cm)
            R_ = RULE_ROW[tf]
            if side > 0:
                put(sg, rr, 14, f"=ROUND(L{rr}+{RM}$D${R_}*M{rr},2)", fmt=RS_)
                put(sg, rr, 15, f"=ROUND(L{rr}+{RM}$C${R_}*M{rr},2)", fmt=RS_)
                put(sg, rr, 16, f"=ROUND(L{rr}-{RM}$B${R_}*M{rr},2)", fmt=RS_)
            elif side < 0:
                put(sg, rr, 14, f"=ROUND(L{rr}-{RM}$D${R_}*M{rr},2)", fmt=RS_)
                put(sg, rr, 15, f"=ROUND(L{rr}-{RM}$C${R_}*M{rr},2)", fmt=RS_)
                put(sg, rr, 16, f"=ROUND(L{rr}+{RM}$B${R_}*M{rr},2)", fmt=RS_)
            if side != 0:
                put(sg, rr, 17, f"=ABS(O{rr}-L{rr})/L{rr}", fmt=PCT)
                put(sg, rr, 18, f"=ABS(L{rr}-P{rr})/L{rr}", fmt=PCT)
            for j in (7, 8, 9):
                sg.cell(rr, j).number_format = PCT
            sg.cell(rr, 10).number_format = PCT_S2
            sg.cell(rr, 11).number_format = PTS
            for j in (12, 13):
                sg.cell(rr, j).number_format = RS_
            for j in (22, 23, 25):
                sg.cell(rr, j).alignment = wrap_top
            rr += 1
    last = rr - 1
    ctx["signals_rows"] = SIG_ROWS
    cf_text(sg, f"F2:F{last}", TYPE_FILL)
    cf_text(sg, f"X2:X{last}", RES_FILL)
    sg.freeze_panes = "E2"
    sg.auto_filter.ref = f"A1:Y{last}"
    widths(sg, [13, 28, 18, 22, 12, 22, 11, 11, 11, 11, 10, 11, 9, 12, 12, 11, 9, 9, 8, 8, 10, 46, 42, 12, 90])

    # =================================================================================================================
    # By Stock
    # =================================================================================================================
    bs = wb.create_sheet("By Stock")
    header(bs, 1, ["Symbol", "Company", "Sector", "Current price Rs", "Short-term call", "Short-term buy confidence",
                   "Mid-term call", "Mid-term buy confidence", "Long-term call", "Long-term buy confidence",
                   "Short-term trend", "Mid-term trend", "Long-term trend", "RS rating (1-99)", "% from 52-wk high",
                   "Active chart patterns (last 10 sessions)", "Results / news check"], 48)
    rr = 2
    for sym in order:
        x = CUR.loc[sym] if sym in CUR.index else None
        vals = [sym, UNI.at[sym, "company"], str(UNI.at[sym, "industry"]).title().replace("It", "IT")]
        if x is None or (sym, "ST") not in Dk:
            vals += [None, "Not rated", None, "Not rated", None, "Not rated", None, "", "", "", None, None, "", ""]
        else:
            vals.append(round(float(x["_close"]), 2))
            for tf in ("ST", "MT", "LT"):
                d = Dk[(sym, tf)]
                vals += [d["type"], d.get("conf_buy") if d["eligible"] else None]
            news = f"{RESEARCH[sym]['verdict']}: {RESEARCH[sym]['note']}" if sym in RESEARCH else "Not checked"
            vals += [trend(x, "ST"), trend(x, "MT"), trend(x, "LT"),
                     None if pd.isna(x.get("xs_rs_rating")) else int(x["xs_rs_rating"]),
                     None if pd.isna(x.get("dist_52w_high")) else float(x["dist_52w_high"]) / 100,
                     active_patterns(sym, 3), news]
        for j, val in enumerate(vals, 1):
            put(bs, rr, j, val)
        bs.cell(rr, 4).number_format = RS_
        for j in (6, 8, 10, 15):
            bs.cell(rr, j).number_format = PCT
        bs.cell(rr, 16).alignment = wrap_top
        bs.cell(rr, 17).alignment = wrap_top
        rr += 1
    for col in "EGI":
        cf_text(bs, f"{col}2:{col}{rr - 1}", TYPE_FILL)
    bs.freeze_panes = "D2"
    bs.auto_filter.ref = f"A1:Q{rr - 1}"
    widths(bs, [13, 28, 18, 11, 22, 11, 22, 11, 22, 11, 10, 10, 10, 9, 10, 44, 50])

    # =================================================================================================================
    # Top Picks
    # =================================================================================================================
    tp = wb.create_sheet("Top Picks")
    tp["A1"] = "Every buy call per timeframe, the strongest 'avoid' flags, and the latest results where checked"
    tp["A1"].font = f_h2
    tp["A2"] = ("Buys: Strong Buy first, then by expected return. " + (ctx.get("research_note", "") + " " if RESEARCH else "")
                + "Levels use the rules on 'Read Me'; recheck the price at the open.")
    tp["A2"].font = f_sub
    rr = 4
    TR = {"Strong Buy": 0, "Buy": 1, "Accumulate (buy on dips)": 1}
    for tf, lab, sm, tm, cm, hold in TFS:
        S = D[(D.tf == tf) & (D.side > 0)].copy()
        S["tr"] = S["type"].map(TR)
        S = S.sort_values(["tr", "exp_ret_buy", "pct_buy"], ascending=[True, False, False])
        tp.cell(rr, 1, f"{lab} - all {len(S)} buy calls").font = f_h2
        rr += 1
        header(tp, rr, ["#", "Symbol", "Company", "Call", "Buy confidence", "Average stock", "Chance of profit",
                        "Expected return / trade", "Entry up to Rs", "Target Rs", "Stop Rs", "Max days", "Why",
                        "Results check", "Results & news", "Source"], 40)
        rr += 1
        for k, (_, d) in enumerate(S.iterrows(), 1):
            sym = d["symbol"]
            res = RESEARCH.get(sym)
            why = "; ".join(t for t in ([f"AI model top {top_pct(d['pct_buy'])}%"] if d.get("long_model") else [])
                            + [p[0] + f" ({p[6]:%d-%b})" for p in sig_list(d, "pos_signals")[:2]])
            vals = [k, sym, UNI.at[sym, "company"], d["type"], d["conf_buy"], d["base_hit_buy"], d["win_buy"],
                    d["exp_ret_buy"], round(d["entry"], 2), round(d["target"], 2), round(d["stop"], 2), hold, why,
                    res["verdict"] if res else "Not checked",
                    (f"{res['quarter']} - revenue: {res['revenue']}; profit: {res['profit']}. {res['note']}"
                     if res else "Not researched"), res["url"] if res else ""]
            for j, val in enumerate(vals, 1):
                put(tp, rr, j, val, wrap=True)
            for j in (5, 6, 7):
                tp.cell(rr, j).number_format = PCT
            tp.cell(rr, 8).number_format = PCT_S2
            for j in (9, 10, 11):
                tp.cell(rr, j).number_format = RS_
            if res and res.get("url"):
                tp.cell(rr, 16).hyperlink = res["url"]
                tp.cell(rr, 16).font = f_link
            tp.row_dimensions[rr].height = 62
            rr += 1
        rr += 1
        A = D[(D.tf == tf) & D["type"].isin(["Avoid (likely underperformer)", "Reduce / Avoid"])].copy()
        A["big"] = A["symbol"].map(lambda s_: float(CUR.loc[s_, "log_turnover_20"]) if s_ in CUR.index else 0)
        A = A.sort_values(["pct_short", "big"], ascending=[False, False]).head(10)
        tp.cell(rr, 1, f"{lab} - strongest avoid / reduce flags").font = f_h2
        rr += 1
        header(tp, rr, ["#", "Symbol", "Company", "Call", "Short-side rank", "Lag vs average stock (hist.)",
                        "Warning signals", "", "", "", "", "", "", "", "", ""], 40)
        rr += 1
        for k, (_, d) in enumerate(A.iterrows(), 1):
            warn = "; ".join(p[0] + f" ({p[6]:%d-%b})" for p in sig_list(d, "neg_signals")[:2]) or "AI model flag"
            vals = [k, d["symbol"], UNI.at[d["symbol"], "company"], d["type"], f"top {top_pct(d['pct_short'])}%",
                    -max(float(d.get("model_ex_short") or 0), 0), warn]
            for j, val in enumerate(vals, 1):
                put(tp, rr, j, val, wrap=True)
            tp.cell(rr, 6).number_format = PCT_S2
            tp.merge_cells(start_row=rr, start_column=7, end_row=rr, end_column=12)
            rr += 1
        rr += 2
    cf_text(tp, f"D4:D{rr}", TYPE_FILL)
    cf_text(tp, f"N4:N{rr}", RES_FILL)
    tp.freeze_panes = "D4"
    widths(tp, [4, 13, 26, 25, 11, 11, 11, 12, 12, 11, 11, 8, 34, 11, 60, 30])

    # =================================================================================================================
    # Pattern Report Card (static research)
    # =================================================================================================================
    pr = wb.create_sheet("Pattern Report Card")
    pr["A1"] = "Report card: does each pattern / signal beat the average stock? (NSE, Nifty 500 stocks, 2017-2026)"
    pr["A1"].font = f_h2
    pr["A2"] = ("Excess = result in the pattern's own direction minus the same-day result of an average eligible stock. "
                "Verdict needs |t| >= 3 (week-clustered) and the same sign in 2017-21 and 2022-26. 'Works in reverse' = the "
                "opposite trade would have worked. Pattern target = measured-move target reached before the pattern's own stop "
                "(chart patterns only). Caution: bearish results look worse than reality because stocks that collapsed left "
                "today's index; 'Survivorship illusion' marks crash-rebound effects that vanish on the full NSE universe.")
    pr["A2"].font = f_sub
    pr["A2"].alignment = wrap_top
    pr.merge_cells("A2:Z2")
    pr.row_dimensions[2].height = 44
    header(pr, 4, ["Source", "Family", "Pattern / signal", "Direction", "Events", "Stocks", "Events since 2024",
                   "2-wk trade: hit %", "2-wk: average stock hit %", "2-wk: excess return / trade", "2-wk: t-stat",
                   "2-wk verdict", "3-mth trade: hit %", "3-mth: average stock hit %", "3-mth: excess return / trade",
                   "3-mth: t-stat", "3-mth verdict", "1-yr: excess return / trade", "1-yr: t-stat", "1-yr verdict",
                   "20-day excess return", "3-mth excess in rising mkts", "3-mth excess in falling mkts",
                   "Pattern target hit %", "Median days to pattern target", "Note"], 56)
    P = RC[RC.level == "pattern"].copy()
    P["src_order"] = P["source"].map({"Chart pattern": 0, "Indicator signal": 1, "Candlestick": 2})
    P = P.sort_values(["src_order", "MT_xnet_t"], ascending=[True, False])
    rr = 5
    for _, x in P.iterrows():
        pn = x["pattern"]
        vals = [x["source"], x["family"], pn, "Bullish" if x["dir"] > 0 else "Bearish", int(x["n"]), int(x["stocks"]),
                int(x["since_2024"]), x["ST_hit"] / 100, x["ST_base"] / 100, x["ST_xnet_mean"] / 100, x["ST_xnet_t"],
                disp_verdict(pn, x["ST_verdict"]), x["MT_hit"] / 100, x["MT_base"] / 100, x["MT_xnet_mean"] / 100,
                x["MT_xnet_t"], disp_verdict(pn, x["MT_verdict"]),
                x["LT_xnet_mean"] / 100, x["LT_xnet_t"], disp_verdict(pn, x["LT_verdict"]), x["x20_mean"] / 100,
                x["MT_xnet_riskon"] / 100, x["MT_xnet_riskoff"] / 100,
                (x["pattern_target_hit"] / 100) if pd.notna(x.get("pattern_target_hit")) else None,
                x.get("median_bars_to_target") if pd.notna(x.get("median_bars_to_target", np.nan)) else None,
                PATTERN_NOTE.get(pn, "")]
        for j, val in enumerate(vals, 1):
            put(pr, rr, j, val, wrap=j == 26)
        for j in (5, 6, 7):
            pr.cell(rr, j).number_format = "#,##0"
        for j in (8, 9, 13, 14, 24):
            pr.cell(rr, j).number_format = PCT
        for j in (10, 15, 18, 21, 22, 23):
            pr.cell(rr, j).number_format = PCT_S2
        for j in (11, 16, 19):
            pr.cell(rr, j).number_format = "0.0"
        rr += 1
    for col in "LQT":
        cf_text(pr, f"{col}5:{col}{rr}", VERDICT_FILL)
    pr.freeze_panes = "D5"
    pr.auto_filter.ref = f"A4:Z{rr - 1}"
    widths(pr, [14, 20, 44, 9, 9, 7, 8, 8, 9, 10, 7, 18, 8, 9, 10, 7, 18, 10, 7, 18, 9, 10, 10, 9, 9, 48])

    # =================================================================================================================
    # Active Patterns
    # =================================================================================================================
    ap = wb.create_sheet("Active Patterns")
    ap["A1"] = f"Chart patterns completed in the last 10 sessions ({recent_from:%d-%b} to {dates[-1]:%d-%b-%Y})"
    ap["A1"].font = f_h2
    ap["A2"] = ("Swing size: minor / intermediate / major = the swing engine's scale; 'bar' = candle- or bar-based pattern. "
                "Verdicts come from the report card - most patterns have no reliable edge on their own.")
    ap["A2"].font = f_sub
    header(ap, 4, ["Date", "Symbol", "Company", "Pattern", "Direction", "Swing size", "Close on signal day Rs",
                   "Trigger level Rs", "Pattern target Rs", "Pattern stop Rs", "Historical verdict (3-mth)",
                   "Historical verdict (2-wk)", "Liquid enough to rate"], 40)
    rr = 5
    E2 = (EV[~EV["family"].isin(["Range contraction"])].sort_values(["date", "symbol"], ascending=[False, True])
          if len(EV) else EV)
    for _, e in E2.iterrows():
        sym = e["symbol"]
        vals = [e["date"].to_pydatetime().date(), sym, UNI.at[sym, "company"] if sym in UNI.index else "", e["pattern"],
                "Bullish" if e["dir"] > 0 else "Bearish", e["scale"], e["close"], e["level"], e["target"], e["stop"],
                disp_verdict(e["pattern"], RCP.at[e["pattern"], "MT_verdict"]) if e["pattern"] in RCP.index else "",
                disp_verdict(e["pattern"], RCP.at[e["pattern"], "ST_verdict"]) if e["pattern"] in RCP.index else "",
                "Yes" if bool(e.get("eligible", False)) else "No"]
        for j, val in enumerate(vals, 1):
            put(ap, rr, j, val)
        ap.cell(rr, 1).number_format = "dd-mmm-yy"
        for j in (7, 8, 9, 10):
            ap.cell(rr, j).number_format = RS_
        rr += 1
    cf_text(ap, f"K5:L{max(rr, 6)}", VERDICT_FILL)
    ap.freeze_panes = "C5"
    ap.auto_filter.ref = f"A4:M{max(rr - 1, 5)}"
    widths(ap, [11, 13, 28, 44, 9, 12, 11, 11, 11, 11, 22, 22, 9])

    # =================================================================================================================
    # Model Validation (static research + today's calibration)
    # =================================================================================================================
    mv = wb.create_sheet("Model Validation")
    mv["A1"] = "Walk-forward validation - every number below is out of sample (the model never saw the year it predicted)"
    mv["A1"].font = f_h2
    lab_tf = {"ST": "Short term (2 weeks)", "MT": "Mid term (3 months)", "LT": "Long term (1 year)"}
    rr = 3
    mv.cell(rr, 1, "1. How well each model ranks stocks against each other on the same day").font = f_bold
    rr += 1
    header(mv, rr, ["Model", "Timeframe", "Side", "Daily rank correlation", "t-stat", "Years positive",
                    "Top 5% vs same-day avg", "Top 15% vs same-day avg", "Bottom 20% vs same-day avg"], 44)
    rr += 1
    for _, x in VAR.iterrows():
        vals = [x["model"], lab_tf[x["tf"]], "Buy" if x["side"] == "L" else "Short", x["daily_ic"], x["ic_t"],
                x["years_positive"], x["top5_excess"] / 100, x["top15_excess"] / 100, x["bottom20_excess"] / 100]
        for j, val in enumerate(vals, 1):
            put(mv, rr, j, val, fill=fill_soft if x["model"].startswith("Absolute") else None)
        mv.cell(rr, 4).number_format = "0.000"
        mv.cell(rr, 5).number_format = "0.0"
        for j in (7, 8, 9):
            mv.cell(rr, j).number_format = PCT_S2
        rr += 1
    c = mv.cell(rr, 1, "Three model designs were tested; the absolute model (shaded) is used because its top-ranked buys "
                       "held up best. Short-side ranks (spotting laggards) are strong in every design.")
    c.font = f_note
    rr += 2
    mv.cell(rr, 1, "2. Rank buckets by market type - excess return per trade vs the same-day average stock").font = f_bold
    rr += 1
    header(mv, rr, ["Timeframe", "Side", "Market", "Bottom 20%", "20-50%", "50-70%", "70-85%", "85-95%", "Top 5%"], 30)
    rr += 1
    for _, x in BER.iterrows():
        vals = [lab_tf[x["tf"]], "Buy" if x["side"] == "L" else "Short", x["regime"]] + \
               [x[k] / 100 for k in ("Bottom 20%", "20-50%", "50-70%", "70-85%", "85-95%", "Top 5%")]
        for j, val in enumerate(vals, 1):
            put(mv, rr, j, val, fill=fill_soft if RW in str(x["regime"]).lower() else None)
        for j in range(4, 10):
            mv.cell(rr, j).number_format = PCT_S2
        rr += 1
    c = mv.cell(rr, 1, "Read across a row: a working ranking rises from left to right. Rows for today's market type are "
                       "shaded.")
    c.font = f_note
    rr += 2
    mv.cell(rr, 1, f"3. Today's calibration ({RW}-market history blended with all markets, survivorship-corrected for buys)"
            ).font = f_bold
    rr += 1
    header(mv, rr, ["Timeframe", "Side", "Rank bucket", "Trades (all)", f"Trades ({RW} mkts)", "Hit target",
                    "Chance of profit", "Avg result / trade", "Excess vs avg (all mkts)", f"Excess vs avg ({RW} mkts)",
                    "Counts as edge"], 44)
    rr += 1
    for (tf, sd), t in BT["tabs"].items():
        for _, x in t["table"].iterrows():
            vals = [lab_tf[tf], "Buy" if sd == "L" else "Short", x["label"], int(x["n_all"]), int(x["n_regime"]), x["y"],
                    x["win"], x["ret"], x["ex_ret_all"], x["ex_ret_regime"], "Yes" if x["edge_ok"] else "No"]
            for j, val in enumerate(vals, 1):
                put(mv, rr, j, val)
            for j in (4, 5):
                mv.cell(rr, j).number_format = "#,##0"
            for j in (6, 7):
                mv.cell(rr, j).number_format = PCT
            for j in (8, 9, 10):
                mv.cell(rr, j).number_format = PCT_S2
            rr += 1
    rr += 1
    mv.cell(rr, 1, "4. AUC by test year (0.5 = coin toss; pooled across days, so it also reflects market timing)").font = f_bold
    rr += 1
    yrs = sorted(FOLD["year"].unique())
    header(mv, rr, ["Timeframe", "Side"] + [str(y) for y in yrs], 30)
    rr += 1
    for (tf, sd), g in FOLD.groupby(["tf", "side"], sort=False):
        put(mv, rr, 1, lab_tf[tf])
        put(mv, rr, 2, "Buy" if sd == "L" else "Short")
        for j, y in enumerate(yrs, 3):
            val = g.loc[g.year == y, "auc"]
            put(mv, rr, j, float(val.iloc[0]) if len(val) else None, fmt="0.000")
        rr += 1
    rr += 1
    mv.cell(rr, 1, "5. Survivorship check - same trade rules on the full NSE universe as it stood each day").font = f_bold
    rr += 1
    header(mv, rr, ["Timeframe", "Universe", "Trades", "Hit target", "Hit stop", "Avg result per trade"], 30)
    rr += 1
    for _, x in SURV.iterrows():
        for j, val in enumerate([lab_tf[x["tf"]], x["universe"], int(x["trades"]), x["hit"] / 100, x["stop"] / 100,
                                 x["avg_net"] / 100], 1):
            put(mv, rr, j, val)
        mv.cell(rr, 3).number_format = "#,##0"
        mv.cell(rr, 4).number_format = PCT
        mv.cell(rr, 5).number_format = PCT
        mv.cell(rr, 6).number_format = PCT_S2
        rr += 1
    rr += 1
    mv.cell(rr, 1, "6. 'Buy after a >=15% two-day crash' - today's list vs the full universe").font = f_bold
    rr += 1
    header(mv, rr, ["Universe", "Crash events", "60-day excess (mean)", "60-day excess (median)", "1-year excess (mean)",
                    "1-year excess (median)", "Share in stocks not in today's list"], 44)
    rr += 1
    for _, x in CRASH.iterrows():
        vals = [x["universe"], int(x["events"]), x["x60_mean"] / 100, x["x60_median"] / 100, x["x250_mean"] / 100,
                x["x250_median"] / 100, x["share_not_in_today_list"]]
        for j, val in enumerate(vals, 1):
            put(mv, rr, j, val)
        for j in (3, 4, 5, 6):
            mv.cell(rr, j).number_format = PCT_S2
        mv.cell(rr, 7).number_format = "0%"
        rr += 1
    widths(mv, [26, 34, 22, 14, 14, 13, 13, 13, 13, 13, 11])

    # =================================================================================================================
    # Indicator Catalogue
    # =================================================================================================================
    ic = wb.create_sheet("Indicator Catalogue")
    ic["A1"] = "Everything the analysis looks at"
    ic["A1"].font = f_h2
    CAT = catalogue(feat_cols, set(FEATS), RC)
    header(ic, 3, list(CAT.columns), 30)
    for i, row in enumerate(CAT.itertuples(index=False), 4):
        for j, val in enumerate(row, 1):
            put(ic, i, j, val)
    ic.auto_filter.ref = f"A3:E{len(CAT) + 3}"
    ic.freeze_panes = "A4"
    widths(ic, [22, 26, 58, 60, 16])

    # =================================================================================================================
    # Charts
    # =================================================================================================================
    if CHARTS:
        ch = wb.create_sheet("Charts")
        ch["A1"] = "Top picks: last 160 sessions, the patterns the engine found, and the trade levels"
        ch["A1"].font = f_h2
        row = 3
        for png in CHARTS:
            img = XLImage(str(png))
            img.width, img.height = 900, 497
            ch.add_image(img, f"A{row}")
            row += 27

    # =================================================================================================================
    # Success Rate (last sheet)
    # =================================================================================================================
    success_sheet(wb, TRACK, ASOF)
    wb.active = 0
    outfile = Path(outfile)
    outfile.parent.mkdir(parents=True, exist_ok=True)
    wb.save(outfile)
    return outfile


def success_sheet(wb, TRACK: dict, asof_label: str):
    ws = wb.create_sheet("Success Rate")
    hd = TRACK.get("headline", {}) or {}
    T = TRACK.get("table", pd.DataFrame())
    ws["A1"] = "Success rate of the AI calls - every Buy / Avoid call since tracking began, checked against real prices"
    ws["A1"].font = f_title
    since = f"Tracking since {hd.get('first_date', '-')}; prices to {asof_label}." if hd else "No calls tracked yet."
    ws["A2"] = since
    ws["A2"].font = f_sub
    ws["A3"] = ("Buy call = paper trade entered at the next session's open only if within the entry limit; closed at target, "
                "stop or after the maximum days; result after 0.3% costs; success = the trade made money. "
                "Avoid call = success if the stock did worse than the average Nifty 500 stock over the same period. "
                "A call repeated while the first one is still running is counted once.")
    ws["A3"].font = f_sub
    ws["A3"].alignment = wrap_top
    ws.merge_cells("A3:N3")
    ws.row_dimensions[3].height = 42
    # ---- headline tiles ------------------------------------------------------------------------------------------
    b, a = hd.get("buy", {}) or {}, hd.get("avoid", {}) or {}
    tiles = [
        ("OVERALL SUCCESS RATE", fmt_pct(hd.get("success_rate", np.nan)),
         f"{hd.get('closed', 0)} completed calls" + (f" | finished groups: {fmt_pct(hd.get('success_rate_mature'))} "
                                                     f"of {hd.get('mature', 0)}" if hd.get("mature") else "")),
        ("BUY CALLS - MADE MONEY", fmt_pct(b.get("success_rate", np.nan)), f"{b.get('closed', 0)} closed trades"),
        ("BUY CALLS - HIT TARGET", fmt_pct(b.get("target_hit_rate", np.nan)),
         f"AI confidence said {fmt_pct(b.get('avg_ai_conf', np.nan))}"),
        ("BUY CALLS - AVG RESULT", "-" if not np.isfinite(b.get("avg_result_pct", np.nan)) else
         f"{b['avg_result_pct']:+.2f}%", "per trade after costs"),
        ("AVOID CALLS - CORRECT", fmt_pct(a.get("success_rate", np.nan)), f"{a.get('closed', 0)} completed"),
        ("STILL RUNNING", str(hd.get("running", 0)), f"{hd.get('not_triggered', 0)} buys not triggered"),
    ]
    r0 = 5
    for k, (lab, val, sub) in enumerate(tiles):
        c0 = 1 + 2 * k
        for rr_, v, f in ((r0, lab, f_tile), (r0 + 1, val, f_big), (r0 + 2, sub, f_note)):
            ws.merge_cells(start_row=rr_, start_column=c0, end_row=rr_, end_column=c0 + 1)
            cell = ws.cell(rr_, c0, v)
            cell.font, cell.alignment, cell.fill = f, center, fill_tile
            ws.cell(rr_, c0 + 1).fill = fill_tile
    ws.row_dimensions[r0].height = 20
    ws.row_dimensions[r0 + 1].height = 36
    ws.row_dimensions[r0 + 2].height = 18
    r = r0 + 4
    c = ws.cell(r, 1, "Reading the numbers: stops sit closer than targets, so losing trades finish first - the success "
                      "rate starts low and settles once trades reach their time exit (10 / 60 / 250 sessions). 'Finished "
                      "groups' counts only days whose calls have ALL completed (no early-close bias). For the model's own "
                      "promise compare 'Buy calls - hit target' with the AI confidence shown under it.")
    c.font, c.alignment = f_note, wrap_top
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
    ws.row_dimensions[r].height = 26
    r += 2
    if not hd.get("closed"):
        c = ws.cell(r, 1, "No call has completed yet. 2-week calls finish after 10 sessions, 3-month calls after 60 and "
                          "1-year calls after 250; the numbers above fill in automatically on each daily run.")
        c.font = Font(name=F, size=10, bold=True, color="92400E")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
        r += 2
    # ---- breakdown -----------------------------------------------------------------------------------------------
    ws.cell(r, 1, "By call type and timeframe").font = f_h2
    r += 1
    cols = ["Call", "Timeframe", "Calls", "Completed", "Running", "Not triggered", "Success rate", "Target hit rate",
            "AI confidence (promised)", "Avg result / trade", "Avg vs average stock (avoid)",
            "Success rate - finished groups only", "Calls in finished groups"]
    header(ws, r, cols, 40)
    r += 1
    G = TRACK.get("by_group", pd.DataFrame())
    first_g = r
    for _, g in (G.iterrows() if len(G) else []):
        vals = [g["Call"], g["Timeframe"], g["calls"], g["closed"], g["running"], g["not_triggered"], g["success_rate"],
                g["target_hit_rate"], g["avg_ai_conf"], None if not np.isfinite(g["avg_result_pct"]) else
                g["avg_result_pct"] / 100, None if not np.isfinite(g["avg_excess_pct"]) else g["avg_excess_pct"] / 100,
                g["success_rate_mature"], g["mature"]]
        for j, v in enumerate(vals, 1):
            put(ws, r, j, v)
        for j in (7, 8, 9, 12):
            ws.cell(r, j).number_format = PCT
        for j in (10, 11):
            ws.cell(r, j).number_format = PCT_S2
        r += 1
    if r > first_g:
        cf_text(ws, f"A{first_g}:A{r - 1}", TYPE_FILL)
    r += 1
    ws.cell(r, 1, "By month of the call").font = f_h2
    r += 1
    header(ws, r, ["Month", "", "Calls", "Completed", "Running", "Not triggered", "Success rate", "Target hit rate",
                   "AI confidence (promised)", "Avg result / trade", "Avg vs average stock (avoid)"], 40)
    r += 1
    Mo = TRACK.get("by_month", pd.DataFrame())
    for _, g in (Mo.iterrows() if len(Mo) else []):
        vals = [g["Month"], "", g["calls"], g["closed"], g["running"], g["not_triggered"], g["success_rate"],
                g["target_hit_rate"], g["avg_ai_conf"], None if not np.isfinite(g["avg_result_pct"]) else
                g["avg_result_pct"] / 100, None if not np.isfinite(g["avg_excess_pct"]) else g["avg_excess_pct"] / 100]
        for j, v in enumerate(vals, 1):
            put(ws, r, j, v)
        for j in (7, 8, 9):
            ws.cell(r, j).number_format = PCT
        for j in (10, 11):
            ws.cell(r, j).number_format = PCT_S2
        r += 1
    # ---- full log ---------------------------------------------------------------------------------------------------
    r += 1
    MAX_ROWS = 5000
    shown = T.head(MAX_ROWS) if len(T) else T
    title = "Every call (newest first)" if len(T) <= MAX_ROWS else \
        f"Latest {MAX_ROWS:,} of {len(T):,} calls (newest first) - full history in output/recommendation_history.csv"
    ws.cell(r, 1, title).font = f_h2
    r += 1
    log_cols = ["Call date", "Symbol", "Company", "Timeframe", "Call", "AI confidence", "Entry limit Rs", "Target Rs",
                "Stop Rs", "Status", "Result", "Entry date", "Entry Rs", "Exit date", "Exit Rs", "Result % (after costs)",
                "Sessions", "Stock vs average stock", "Repeated (days)", "Note", "Why"]
    header(ws, r, log_cols, 40)
    hdr_row = r
    r += 1
    tf_lab = {"ST": "2 weeks", "MT": "3 months", "LT": "1 year"}
    for _, x in (shown.iterrows() if len(shown) else []):
        vals = [pd.Timestamp(x["signal_date"]).date(), x["symbol"], x["company"], tf_lab.get(x["timeframe"], x["timeframe"]),
                x["call_type"], x["ai_conf"], x["entry_limit"], x["target"], x["stop"], x["status"], x["outcome"],
                pd.Timestamp(x["entry_date"]).date() if isinstance(x["entry_date"], str) else None, x["entry_price"],
                pd.Timestamp(x["exit_date"]).date() if isinstance(x["exit_date"], str) else None, x["exit_price"],
                None if pd.isna(x["result_pct"]) else x["result_pct"] / 100,
                None if pd.isna(x["days_held"]) else int(x["days_held"]),
                None if pd.isna(x["excess_pct"]) else x["excess_pct"] / 100, int(x["repeat_count"] or 0), x["note"],
                x["evidence"]]
        for j, v in enumerate(vals, 1):
            put(ws, r, j, v)
        for j in (1, 12, 14):
            ws.cell(r, j).number_format = "dd-mmm-yy"
        ws.cell(r, 6).number_format = PCT
        for j in (7, 8, 9, 13, 15):
            ws.cell(r, j).number_format = RS_
        for j in (16, 18):
            ws.cell(r, j).number_format = PCT_S2
        r += 1
    if r > hdr_row + 1:
        cf_text(ws, f"J{hdr_row + 1}:J{r - 1}", STATUS_FILL)
        cf_text(ws, f"K{hdr_row + 1}:K{r - 1}", OUTCOME_FILL)
        cf_text(ws, f"E{hdr_row + 1}:E{r - 1}", TYPE_FILL)
        ws.auto_filter.ref = f"A{hdr_row}:U{r - 1}"
    widths(ws, [11, 13, 26, 10, 22, 11, 11, 11, 11, 13, 11, 12, 11, 11, 11, 12, 9, 12, 10, 44, 50])
